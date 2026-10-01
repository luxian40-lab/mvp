from django.shortcuts import render
from django.contrib.admin.views.decorators import staff_member_required
from django.http import HttpResponse, JsonResponse
from django.conf import settings
from datetime import datetime, timedelta
import json
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.utils import timezone

logger = logging.getLogger(__name__)

from ..models import Campana, Estudiante, WhatsappLog, EnvioLog, Cliente, Curso, ProgresoEstudiante, ModuloCompletado
from ..models_extras import ArchivoModulo, GrupoEstudiantes


def _construir_dashboard_unificado_contexto(request, incluir_detalle=True):
    """Construye contexto y payload JSON del dashboard unificado usando filtros consistentes."""
    from datetime import datetime, timedelta
    from django.db.models import Avg, Count, Q
    from django.db.models.functions import TruncDate
    from django.utils import timezone
    import json

    def _to_int(value):
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    def _to_date(value):
        if not value:
            return None
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except (TypeError, ValueError):
            return None

    cliente_id = _to_int((request.GET.get('cliente') or '').strip())
    curso_id = _to_int((request.GET.get('curso') or '').strip())
    fecha_inicio_raw = (request.GET.get('fecha_inicio') or '').strip()
    fecha_fin_raw = (request.GET.get('fecha_fin') or '').strip()
    municipio_filtro = (request.GET.get('municipio') or '').strip()
    tab_raw = (request.GET.get('tab') or 'executive').strip().lower()
    from core.domains.dashboard import resolve_dashboard_tab, resolve_learning_section

    tab_actual = resolve_dashboard_tab(tab_raw)
    learning_section = resolve_learning_section(tab_raw, request.GET.get('section'))
    grupo_id = _to_int((request.GET.get('grupo_id') or request.GET.get('grupo') or '').strip())
    modulo_hasta_numero = _to_int((request.GET.get('modulo_hasta') or '').strip())

    fecha_inicio_dt = _to_date(fecha_inicio_raw)
    fecha_fin_dt = _to_date(fecha_fin_raw)
    if fecha_inicio_dt and fecha_fin_dt and fecha_inicio_dt > fecha_fin_dt:
        fecha_inicio_dt, fecha_fin_dt = fecha_fin_dt, fecha_inicio_dt

    fecha_inicio = fecha_inicio_dt.isoformat() if fecha_inicio_dt else ''
    fecha_fin = fecha_fin_dt.isoformat() if fecha_fin_dt else ''

    # Retención: contexto liviano (evita recalcular Executive/Learning completo).
    if tab_actual == 'retencion':
        from portal.retencion_service import analitica_retencion_portal

        clientes_all = Cliente.objects.all().order_by('nombre')
        cursos_all = Curso.objects.filter(activo=True).order_by('orden', 'nombre')
        if cliente_id:
            cursos_all = cursos_all.filter(cliente_id=cliente_id)
            if curso_id and not cursos_all.filter(pk=curso_id).exists():
                curso_id = None
        grupos_qs = GrupoEstudiantes.objects.all().order_by('nombre')
        if cliente_id:
            grupos_qs = grupos_qs.filter(cliente_id=cliente_id)

        retencion_data = None
        if cliente_id:
            org = Cliente.objects.filter(pk=cliente_id, activo=True).first()
            if org:
                desde_ret = (request.GET.get('desde') or fecha_inicio or '').strip() or None
                hasta_ret = (request.GET.get('hasta') or fecha_fin or '').strip() or None
                retencion_data = analitica_retencion_portal(
                    org,
                    curso_id=curso_id,
                    grupo_id=grupo_id,
                    desde=desde_ret,
                    hasta=hasta_ret,
                )

        context = {
            'tab_actual': tab_actual,
            'learning_section': learning_section,
            'clientes': Cliente.objects.all().order_by('nombre'),
            'cursos': cursos_all,
            'cursos_retencion': cursos_all,
            'cliente_filtro': cliente_id,
            'curso_filtro': curso_id,
            'fecha_inicio': fecha_inicio,
            'fecha_fin': fecha_fin,
            'municipios': [],
            'municipio_filtro': municipio_filtro,
            'grupos': grupos_qs,
            'grupo_filtro': grupo_id,
            'modulo_hasta_filtro': modulo_hasta_numero,
            'retencion_data': retencion_data,
            'clientes_detalle': [],
            'estudiantes_detalle': [],
            'tickets_soporte': [],
            'eventos_ia_recientes': [],
            'embudo_learning': None,
            'resumen_payload_json': '{}',
            'chart_labels': '[]',
            'chart_values': '[]',
            'chart_ubicaciones_labels': '[]',
            'chart_ubicaciones_values': '[]',
            'chart_tipos_labels': '[]',
            'chart_tipos_values': '[]',
            'total_cursos': 0,
            'total_clientes': 0,
            'total_estudiantes': 0,
            'total_mensajes_whatsapp': 0,
            'mensajes_enviados': 0,
            'mensajes_recibidos': 0,
            'wa_entregados': 0,
            'wa_leidos': 0,
            'wa_en_transito': 0,
            'wa_bot_comercial_sent': 0,
            'wa_bot_comercial_read': 0,
            'total_audios': 0,
            'total_agentes_ia': 0,
            'total_progreso': 0,
            'total_modulos_completados': 0,
            'cursos_completados': 0,
            'total_certificados': 0,
            'total_perfiles_gam': 0,
            'puntos_promedio': 0,
            'top_estudiantes': [],
            'ranking_gamificacion_completo': [],
            'total_prospectos': 0,
            'tasa_completacion': 0,
            'ubicaciones_municipio': [],
            'progreso_por_curso': [],
        }
        return context, {}

    clientes_all = Cliente.objects.all().order_by('nombre')
    # Con organización elegida: solo cursos de esa org (no el catálogo completo).
    cursos_all = Curso.objects.all().order_by('nombre')
    if cliente_id:
        cursos_all = cursos_all.filter(cliente_id=cliente_id)
        if curso_id and not cursos_all.filter(pk=curso_id).exists():
            curso_id = None

    estudiantes_q = Estudiante.objects.filter(activo=True)
    if cliente_id:
        estudiantes_q = estudiantes_q.filter(cliente_id=cliente_id)
    if curso_id:
        estudiantes_q = estudiantes_q.filter(progresos__curso_id=curso_id).distinct()
    if fecha_inicio_dt:
        estudiantes_q = estudiantes_q.filter(fecha_registro__date__gte=fecha_inicio_dt)
    if fecha_fin_dt:
        estudiantes_q = estudiantes_q.filter(fecha_registro__date__lte=fecha_fin_dt)
    if grupo_id:
        estudiantes_q = estudiantes_q.filter(grupos__id=grupo_id).distinct()

    progreso_q = ProgresoEstudiante.objects.select_related('estudiante', 'curso')
    if cliente_id:
        progreso_q = progreso_q.filter(estudiante__cliente_id=cliente_id)
    if curso_id:
        progreso_q = progreso_q.filter(curso_id=curso_id)
    if fecha_inicio_dt:
        progreso_q = progreso_q.filter(fecha_inicio__date__gte=fecha_inicio_dt)
    if fecha_fin_dt:
        progreso_q = progreso_q.filter(fecha_inicio__date__lte=fecha_fin_dt)
    if grupo_id:
        progreso_q = progreso_q.filter(estudiante__grupos__id=grupo_id).distinct()

    modulos_completados_q = ModuloCompletado.objects.select_related('progreso', 'modulo')
    if cliente_id:
        modulos_completados_q = modulos_completados_q.filter(progreso__estudiante__cliente_id=cliente_id)
    if curso_id:
        modulos_completados_q = modulos_completados_q.filter(progreso__curso_id=curso_id)
    if fecha_inicio_dt:
        modulos_completados_q = modulos_completados_q.filter(fecha_completado__date__gte=fecha_inicio_dt)
    if fecha_fin_dt:
        modulos_completados_q = modulos_completados_q.filter(fecha_completado__date__lte=fecha_fin_dt)
    if grupo_id:
        modulos_completados_q = modulos_completados_q.filter(
            progreso__estudiante__grupos__id=grupo_id
        ).distinct()

    whatsapp_q = WhatsappLog.objects.all()
    if fecha_inicio_dt:
        whatsapp_q = whatsapp_q.filter(fecha__date__gte=fecha_inicio_dt)
    if fecha_fin_dt:
        whatsapp_q = whatsapp_q.filter(fecha__date__lte=fecha_fin_dt)

    if cliente_id or curso_id:
        estudiantes_scope = Estudiante.objects.filter(activo=True)
        if cliente_id:
            estudiantes_scope = estudiantes_scope.filter(cliente_id=cliente_id)
        if curso_id:
            estudiantes_scope = estudiantes_scope.filter(progresos__curso_id=curso_id).distinct()
        if grupo_id:
            estudiantes_scope = estudiantes_scope.filter(grupos__id=grupo_id).distinct()

        telefonos_scope = estudiantes_scope.exclude(telefono='').values_list('telefono', flat=True)
        whatsapp_q = whatsapp_q.filter(
            Q(estudiante__in=estudiantes_scope) | Q(telefono__in=telefonos_scope)
        ).distinct()

    total_estudiantes = estudiantes_q.count()

    if cliente_id:
        total_clientes = Cliente.objects.filter(id=cliente_id).count()
    elif curso_id or fecha_inicio_dt or fecha_fin_dt:
        total_clientes = Cliente.objects.filter(estudiantes__in=estudiantes_q).distinct().count()
    else:
        total_clientes = Cliente.objects.count()

    progreso_filter_q = Q(progresoestudiante__id__in=progreso_q.values('id'))
    progreso_por_curso_qs = Curso.objects.all()
    if cliente_id or curso_id or fecha_inicio_dt or fecha_fin_dt:
        progreso_por_curso_qs = progreso_por_curso_qs.filter(progreso_filter_q)

    progreso_por_curso = progreso_por_curso_qs.annotate(
        total_estudiantes=Count('progresoestudiante', filter=progreso_filter_q, distinct=True),
        total_modulos_completados=Count('progresoestudiante__modulos_completados', filter=progreso_filter_q, distinct=True),
        completados=Count(
            'progresoestudiante',
            filter=progreso_filter_q & Q(progresoestudiante__completado=True),
            distinct=True,
        ),
    ).order_by('nombre')

    total_cursos = Curso.objects.count() if not (cliente_id or curso_id or fecha_inicio_dt or fecha_fin_dt) else progreso_por_curso.count()

    total_mensajes_whatsapp = whatsapp_q.count()
    mensajes_enviados = whatsapp_q.filter(tipo='SENT').count()
    mensajes_recibidos = whatsapp_q.filter(tipo='INCOMING').count()
    total_audios = whatsapp_q.filter(es_audio=True).count()
    total_agentes_ia = whatsapp_q.filter(agente_usado__isnull=False).exclude(agente_usado='').count()

    wa_entregados = whatsapp_q.filter(tipo='SENT', estado__iexact='DELIVERED').count()
    wa_leidos = whatsapp_q.filter(tipo='SENT', estado__iexact='READ').count()
    wa_en_transito = whatsapp_q.filter(
        tipo='SENT',
        estado__in=['PENDING', 'QUEUED', 'SENDING', 'pending', 'queued', 'sending'],
    ).count()
    wa_bot_comercial_sent = whatsapp_q.filter(tipo='SENT', agente_usado='BOT_COMERCIAL').count()
    wa_bot_comercial_read = whatsapp_q.filter(
        tipo='SENT', agente_usado='BOT_COMERCIAL', estado__iexact='READ'
    ).count()

    total_progreso = progreso_q.count()
    total_modulos_completados = modulos_completados_q.count()
    cursos_completados = progreso_q.filter(completado=True).count()

    try:
        from ..models_certificados import Certificado
        certificados_q = Certificado.objects.all()
        if cliente_id:
            certificados_q = certificados_q.filter(estudiante__cliente_id=cliente_id)
        if curso_id:
            certificados_q = certificados_q.filter(curso_id=curso_id)
        if fecha_inicio_dt:
            certificados_q = certificados_q.filter(fecha_emision__date__gte=fecha_inicio_dt)
        if fecha_fin_dt:
            certificados_q = certificados_q.filter(fecha_emision__date__lte=fecha_fin_dt)
        total_certificados = certificados_q.count()
    except Exception:
        total_certificados = 0

    total_perfiles_gam = 0
    puntos_promedio = 0
    top_estudiantes = []
    ranking_gamificacion_completo = []
    try:
        from ..gamificacion import PerfilGamificacion
        perfiles_q = PerfilGamificacion.objects.filter(estudiante_id__in=estudiantes_q.values('id'))
        total_perfiles_gam = perfiles_q.count()
        puntos_promedio = perfiles_q.aggregate(avg=Avg('puntos_totales'))['avg'] or 0
        if incluir_detalle:
            rank_qs = (
                perfiles_q.select_related('estudiante', 'estudiante__cliente')
                .order_by('-puntos_totales')
            )
            top_estudiantes = rank_qs[:10]
            ranking_gamificacion_completo = list(rank_qs[:2000])
    except Exception:
        pass

    ubicaciones_municipio = (
        estudiantes_q.exclude(municipio__isnull=True)
        .exclude(municipio='')
        .values('municipio')
        .annotate(total=Count('id'))
        .order_by('-total')[:10]
    )
    chart_ubicaciones_labels = [u['municipio'] for u in ubicaciones_municipio]
    chart_ubicaciones_values = [u['total'] for u in ubicaciones_municipio]

    hoy = fecha_fin_dt or timezone.localdate()
    hace_7_dias = hoy - timedelta(days=6)
    mensajes_por_dia = (
        whatsapp_q.filter(fecha__date__gte=hace_7_dias, fecha__date__lte=hoy)
        .annotate(dia=TruncDate('fecha'))
        .values('dia')
        .annotate(total=Count('id'))
        .order_by('dia')
    )
    mensajes_por_dia_map = {m['dia']: int(m['total']) for m in mensajes_por_dia}

    chart_labels = []
    chart_values = []
    for i in range(7):
        dia = hoy - timedelta(days=6 - i)
        chart_labels.append(dia.strftime('%d/%m'))
        chart_values.append(int(mensajes_por_dia_map.get(dia, 0)))

    tipos_msg = whatsapp_q.values('tipo').annotate(total=Count('id')).order_by('-total')
    chart_tipos_labels = [t['tipo'] or 'Otro' for t in tipos_msg]
    chart_tipos_values = [int(t['total']) for t in tipos_msg]

    try:
        from ..models import ProspectoB2B
        prospectos_q = ProspectoB2B.objects.all()
        if fecha_inicio_dt:
            prospectos_q = prospectos_q.filter(fecha_captura__date__gte=fecha_inicio_dt)
        if fecha_fin_dt:
            prospectos_q = prospectos_q.filter(fecha_captura__date__lte=fecha_fin_dt)
        total_prospectos = prospectos_q.count()
    except Exception:
        total_prospectos = 0

    total_inscripciones = total_progreso
    tasa_completacion = round((cursos_completados / total_inscripciones * 100), 1) if total_inscripciones > 0 else 0

    municipios = list(
        estudiantes_q.exclude(municipio__isnull=True)
        .exclude(municipio='')
        .values_list('municipio', flat=True)
        .distinct()
        .order_by('municipio')
    )

    estudiantes_detalle = []
    clientes_detalle = []
    tickets_soporte = []
    if incluir_detalle:
        est_q = estudiantes_q.select_related('cliente').prefetch_related('grupos')
        if municipio_filtro:
            est_q = est_q.filter(municipio=municipio_filtro)

        est_ids = list(est_q[:200].values_list('id', flat=True))

        puntos_map = {}
        try:
            from ..gamificacion import PerfilGamificacion as PG_detail
            puntos_map = dict(
                PG_detail.objects.filter(estudiante_id__in=est_ids).values_list('estudiante_id', 'puntos_totales')
            )
        except Exception:
            puntos_map = {}

        progresos_rows = (
            progreso_q.filter(estudiante_id__in=est_ids)
            .select_related('estudiante', 'estudiante__cliente', 'curso', 'modulo_actual')
            .prefetch_related('estudiante__grupos')
            .annotate(
                total_mods=Count('curso__modulos', distinct=True),
                mods_comp=Count('modulos_completados', distinct=True),
            )
            .order_by('estudiante__nombre', 'curso__nombre')
        )

        seen_estudiante_sin_progreso = set()
        from core.drip_schedule import avance_sobre_modulos, estudiante_llego_hasta_modulo, modulos_para_metricas

        for progreso in progresos_rows:
            est = progreso.estudiante
            if modulo_hasta_numero is not None and not estudiante_llego_hasta_modulo(progreso, modulo_hasta_numero):
                continue
            seen_estudiante_sin_progreso.add(est.id)
            total_mods = progreso.total_mods or 0
            mods_comp = progreso.mods_comp or 0
            avance = round(mods_comp / total_mods * 100) if total_mods > 0 else 0
            mods_drip = modulos_para_metricas(
                est,
                progreso.curso,
                modulo_hasta_numero=modulo_hasta_numero,
                usar_drip_calendario=modulo_hasta_numero is None,
            )
            comps_drip, total_drip, avance_drip = avance_sobre_modulos(progreso, mods_drip)
            if progreso.completado:
                estado_avance = 'Completado'
                modulo_txt = 'Curso completado'
            elif progreso.modulo_actual_id and progreso.modulo_actual:
                estado_avance = 'En curso'
                m = progreso.modulo_actual
                modulo_txt = f'M{m.numero} · {m.titulo}'
            elif avance > 0:
                estado_avance = 'En curso'
                modulo_txt = f'En curso ({mods_comp}/{total_mods} módulos)'
            else:
                estado_avance = 'Sin avance'
                modulo_txt = 'Sin iniciar'
            grupos_txt = ', '.join(sorted(g.nombre for g in est.grupos.all())) or '-'
            estudiantes_detalle.append({
                'nombre': est.nombre,
                'cedula': est.cedula,
                'telefono': (est.telefono or '').strip() or '-',
                'organizacion': est.cliente.nombre if est.cliente else '-',
                'municipio': est.municipio or '-',
                'curso': progreso.curso.nombre if progreso.curso_id else '-',
                'modulo_actual': modulo_txt,
                'modulos_completados': f'{mods_comp}/{total_mods}' if total_mods else '-',
                'modulos_drip': f'{comps_drip}/{total_drip}' if total_drip else '-',
                'avance': avance,
                'avance_drip': avance_drip,
                'puntos': puntos_map.get(est.id, 0),
                'grupos': grupos_txt,
                'estado_avance': estado_avance,
            })

        for est in est_q.filter(id__in=est_ids).exclude(id__in=seen_estudiante_sin_progreso):
            grupos_txt = ', '.join(sorted(g.nombre for g in est.grupos.all())) or '-'
            estudiantes_detalle.append({
                'nombre': est.nombre,
                'cedula': est.cedula,
                'telefono': (est.telefono or '').strip() or '-',
                'organizacion': est.cliente.nombre if est.cliente else '-',
                'municipio': est.municipio or '-',
                'curso': '-',
                'modulo_actual': '-',
                'modulos_completados': '-',
                'modulos_drip': '-',
                'avance': 0,
                'avance_drip': 0,
                'puntos': puntos_map.get(est.id, 0),
                'grupos': grupos_txt,
                'estado_avance': 'Sin inscripción',
            })

        clientes_iter = clientes_all if not cliente_id else clientes_all.filter(id=cliente_id)
        for c in clientes_iter:
            est_cliente = estudiantes_q.filter(cliente=c)
            n_est = est_cliente.count()
            tels = est_cliente.exclude(telefono='').values_list('telefono', flat=True)
            progreso_cliente = progreso_q.filter(estudiante__cliente=c)

            whatsapp_cliente = whatsapp_q.filter(
                Q(estudiante__cliente=c) | Q(telefono__in=tels)
            ).distinct()

            clientes_detalle.append({
                'nombre': c.nombre,
                'cursos': progreso_cliente.values('curso_id').distinct().count(),
                'estudiantes': n_est,
                'uso_audio': whatsapp_cliente.filter(es_audio=True).count(),
                'uso_ia': whatsapp_cliente.filter(agente_usado__isnull=False).exclude(agente_usado='').count(),
                'cursos_completados': progreso_cliente.filter(completado=True).count(),
            })

        try:
            from ..models import SolicitudSoporte
            tickets_soporte_q = SolicitudSoporte.objects.select_related('estudiante')
            if cliente_id:
                tickets_soporte_q = tickets_soporte_q.filter(estudiante__cliente_id=cliente_id)
            if curso_id:
                tickets_soporte_q = tickets_soporte_q.filter(estudiante__progresos__curso_id=curso_id).distinct()
            if fecha_inicio_dt:
                tickets_soporte_q = tickets_soporte_q.filter(fecha_solicitud__date__gte=fecha_inicio_dt)
            if fecha_fin_dt:
                tickets_soporte_q = tickets_soporte_q.filter(fecha_solicitud__date__lte=fecha_fin_dt)
            tickets_soporte = tickets_soporte_q.order_by('-fecha_solicitud')[:50]
        except Exception:
            tickets_soporte = []

    grupos_qs = GrupoEstudiantes.objects.all().order_by('nombre')
    if cliente_id:
        grupos_qs = grupos_qs.filter(cliente_id=cliente_id)

    resumen_payload = {
        'success': True,
        'generated_at': timezone.now().isoformat(),
        'kpis': {
            'total_cursos': int(total_cursos),
            'total_clientes': int(total_clientes),
            'total_estudiantes': int(total_estudiantes),
            'total_mensajes_whatsapp': int(total_mensajes_whatsapp),
            'mensajes_enviados': int(mensajes_enviados),
            'mensajes_recibidos': int(mensajes_recibidos),
            'wa_entregados': int(wa_entregados),
            'wa_leidos': int(wa_leidos),
            'wa_en_transito': int(wa_en_transito),
            'wa_bot_comercial_sent': int(wa_bot_comercial_sent),
            'wa_bot_comercial_read': int(wa_bot_comercial_read),
            'total_audios': int(total_audios),
            'total_agentes_ia': int(total_agentes_ia),
            'total_progreso': int(total_progreso),
            'total_modulos_completados': int(total_modulos_completados),
            'cursos_completados': int(cursos_completados),
            'total_certificados': int(total_certificados),
            'total_perfiles_gam': int(total_perfiles_gam),
            'puntos_promedio': round(float(puntos_promedio), 1),
            'total_prospectos': int(total_prospectos),
            'tasa_completacion': float(tasa_completacion),
        },
        'chart_mensajes': {
            'labels': chart_labels,
            'values': chart_values,
        },
        'chart_ubicaciones': {
            'labels': chart_ubicaciones_labels,
            'values': chart_ubicaciones_values,
        },
        'chart_tipos': {
            'labels': chart_tipos_labels,
            'values': chart_tipos_values,
        },
    }

    eventos_ia_recientes = []
    try:
        from core.models import EventoIA

        eventos_ia_recientes = list(
            EventoIA.objects.select_related('estudiante', 'curso', 'modulo')
            .order_by('-created_at')[:20]
        )
    except Exception:
        eventos_ia_recientes = []

    context = {
        'total_cursos': total_cursos,
        'total_clientes': total_clientes,
        'total_estudiantes': total_estudiantes,
        'total_mensajes_whatsapp': total_mensajes_whatsapp,
        'mensajes_enviados': mensajes_enviados,
        'mensajes_recibidos': mensajes_recibidos,
        'wa_entregados': wa_entregados,
        'wa_leidos': wa_leidos,
        'wa_en_transito': wa_en_transito,
        'wa_bot_comercial_sent': wa_bot_comercial_sent,
        'wa_bot_comercial_read': wa_bot_comercial_read,
        'total_audios': total_audios,
        'total_agentes_ia': total_agentes_ia,
        'total_progreso': total_progreso,
        'total_modulos_completados': total_modulos_completados,
        'cursos_completados': cursos_completados,
        'total_certificados': total_certificados,
        'total_perfiles_gam': total_perfiles_gam,
        'puntos_promedio': round(puntos_promedio, 1),
        'top_estudiantes': top_estudiantes,
        'ranking_gamificacion_completo': ranking_gamificacion_completo,
        'total_prospectos': total_prospectos,
        'tasa_completacion': tasa_completacion,
        'ubicaciones_municipio': ubicaciones_municipio,
        'progreso_por_curso': progreso_por_curso,
        'chart_labels': json.dumps(chart_labels),
        'chart_values': json.dumps(chart_values),
        'chart_ubicaciones_labels': json.dumps(chart_ubicaciones_labels),
        'chart_ubicaciones_values': json.dumps(chart_ubicaciones_values),
        'chart_tipos_labels': json.dumps(chart_tipos_labels),
        'chart_tipos_values': json.dumps(chart_tipos_values),
        'resumen_payload_json': json.dumps(resumen_payload),
        'clientes': clientes_all,
        'cursos': cursos_all,
        'cliente_filtro': cliente_id,
        'curso_filtro': curso_id,
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'municipios': municipios,
        'municipio_filtro': municipio_filtro,
        'tab_actual': tab_actual,
        'learning_section': learning_section,
        'grupos': grupos_qs,
        'grupo_filtro': grupo_id,
        'modulo_hasta_filtro': modulo_hasta_numero,
        'clientes_detalle': clientes_detalle,
        'estudiantes_detalle': estudiantes_detalle,
        'tickets_soporte': tickets_soporte,
        'eventos_ia_recientes': eventos_ia_recientes,
        'retencion_data': None,
        'cursos_retencion': Curso.objects.none(),
        'embudo_learning': None,
    }

    if learning_section == 'embudo' and curso_id:
        from portal.curso_flujo_service import embudo_posicion_hoy_por_curso

        context['embudo_learning'] = embudo_posicion_hoy_por_curso(
            curso_id=curso_id,
            cliente_id=cliente_id,
            grupo_id=grupo_id,
        )

    context['video_reporte'] = None
    if learning_section == 'videos':
        from core.video_links import reporte_aperturas_video

        context['video_reporte'] = reporte_aperturas_video(
            curso_id=curso_id,
            cliente_id=cliente_id,
            fecha_inicio=fecha_inicio_dt,
            fecha_fin=fecha_fin_dt,
        )

    return context, resumen_payload


# Vista unificada del dashboard admin
@staff_member_required
def dashboard_unificado(request):
    """
    Dashboard profesional unificado de eki.
    Métricas reales: cursos, clientes, estudiantes, certificados, gamificación,
    WhatsApp, IA, y progreso educativo.
    """
    from core.domains.dashboard import resolve_dashboard_tab, resolve_learning_section

    _tab = resolve_dashboard_tab(request.GET.get('tab'))
    _section = resolve_learning_section(_tab, request.GET.get('section'))
    _exportando = (request.GET.get('exportar') or '').strip().lower() == 'excel'
    # Detalle pesado solo con ?detalle=1 o Excel (antes: siempre en Reportes → Analítica lenta).
    _quiere_detalle = (request.GET.get('detalle') or '').strip() in ('1', 'true', 'si', 'sí')
    context, resumen_payload = _construir_dashboard_unificado_contexto(
        request,
        incluir_detalle=_exportando or _quiere_detalle,
    )
    context['detalle_cargado'] = bool(_exportando or _quiere_detalle)

    # --- Excel export (todas las pestañas + datos de gráficos) ---
    if request.GET.get('exportar') == 'excel':
        from analytics.exports import export_dashboard_excel

        return export_dashboard_excel(
            context=context,
            resumen_payload=resumen_payload,
            tab=context.get('tab_actual', 'executive'),
            learning_section=context.get('learning_section', 'reportes'),
        )

    return render(request, 'admin/dashboard.html', context)


@staff_member_required
def dashboard_unificado_resumen_data(request):
    """Endpoint JSON para refresco en tiempo real del Resumen Ejecutivo."""
    _, payload = _construir_dashboard_unificado_contexto(request, incluir_detalle=False)
    return JsonResponse(payload)


@staff_member_required
def bot_comercial_admin_view(request):
    """Vista administrativa para operación del Bot Comercial IA."""
    from datetime import timedelta
    from django.urls import reverse
    from django.utils import timezone as dj_tz

    from ..models import DocumentoRAGComercial, WhatsappLog
    from ..rag_comercial_manager import rag_comercial_manager

    endpoint_path = '/webhook/ia-bot-comercial/'
    endpoint_url = request.build_absolute_uri(endpoint_path)
    cliente_id = int(
        getattr(settings, 'BOT_COMERCIAL_CLIENTE_ID', 0) or 0
    )
    canal_rag = str(getattr(settings, 'BOT_COMERCIAL_RAG_CANAL', 'bot_comercial') or 'bot_comercial')

    docs_qs = DocumentoRAGComercial.objects.filter(cliente_id=cliente_id if cliente_id > 0 else None, canal=canal_rag)
    total_docs = docs_qs.count()
    total_docs_indexados = docs_qs.filter(estado='indexado').count()

    chunks_total = 0
    if rag_comercial_manager.disponible:
        chunks_total = rag_comercial_manager.contar_chunks(cliente_id=cliente_id, canal=canal_rag)

    hace_7 = dj_tz.now() - timedelta(days=7)
    bc_in = WhatsappLog.objects.filter(
        tipo='INCOMING', agente_usado='BOT_COMERCIAL', fecha__gte=hace_7
    ).count()
    bc_out = WhatsappLog.objects.filter(
        tipo='SENT', agente_usado='BOT_COMERCIAL', fecha__gte=hace_7
    ).count()
    bc_read = WhatsappLog.objects.filter(
        tipo='SENT', agente_usado='BOT_COMERCIAL', estado__iexact='READ'
    ).count()
    bc_delivered = WhatsappLog.objects.filter(
        tipo='SENT', agente_usado='BOT_COMERCIAL', estado__iexact='DELIVERED'
    ).count()
    callback_general = request.build_absolute_uri('/webhook/whatsapp/')

    context = {
        'endpoint_path': endpoint_path,
        'endpoint_url': endpoint_url,
        'bot_comercial_whatsapp_number': getattr(settings, 'BOT_COMERCIAL_WHATSAPP_NUMBER', ''),
        'bot_comercial_cliente_id': getattr(settings, 'BOT_COMERCIAL_CLIENTE_ID', ''),
        'bot_comercial_rag_canal': canal_rag,
        'bot_comercial_force_routing': bool(getattr(settings, 'BOT_COMERCIAL_FORCE_ROUTING', False)),
        'bot_comercial_openai_model': getattr(settings, 'BOT_COMERCIAL_OPENAI_MODEL', ''),
        'bot_comercial_vision_model': getattr(settings, 'BOT_COMERCIAL_VISION_MODEL', ''),
        'rag_comercial_disponible': rag_comercial_manager.disponible,
        'rag_comercial_total_docs': total_docs,
        'rag_comercial_docs_indexados': total_docs_indexados,
        'rag_comercial_chunks_total': chunks_total,
        'bot_comercial_incoming_7d': bc_in,
        'bot_comercial_sent_7d': bc_out,
        'bot_comercial_read_total': bc_read,
        'bot_comercial_delivered_total': bc_delivered,
        'memory_turnos': int(getattr(settings, 'BOT_COMERCIAL_MEMORY_TURNOS', 12) or 12),
        'memory_chars': int(getattr(settings, 'BOT_COMERCIAL_MEMORY_MAX_CHARS', 3600) or 3600),
        'twilio_status_callback_configured': bool(
            str(getattr(settings, 'TWILIO_STATUS_CALLBACK_URL', '') or '').strip()
        ),
        'twilio_status_callback_example': callback_general,
        'whatsapp_log_admin_url': reverse('admin:core_whatsapplog_changelist'),
    }
    return render(request, 'admin/bot_comercial.html', context)


@staff_member_required
def dashboard_view(request):
    """Redirige al dashboard de métricas completo (contexto histórico incompleto aquí rompía la plantilla)."""
    from django.shortcuts import redirect
    from django.urls import reverse

    target = reverse('dashboard_metrics')
    qs = request.META.get('QUERY_STRING', '')
    if qs:
        return redirect(f'{target}?{qs}')
    return redirect(target)


# ---------- Vista de instrucciones ----------
@staff_member_required
def instrucciones_view(request):
    """Vista para mostrar el instructivo completo de eki."""
    return render(request, 'admin/instrucciones.html')


# ---------- Vista de importación de prospectos B2B ----------
@staff_member_required
def importar_prospectos(request):
    """Importar prospectos B2B desde archivo Excel.
    Formato: Teléfono | Nombre Contacto | Email | Empresa
    """
    import re
    context = {}

    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')
        if not archivo:
            context['error'] = "Por favor selecciona un archivo Excel"
            return render(request, 'admin/importar_prospectos.html', context)

        try:
            if not archivo.name.endswith(('.xlsx', '.xls')):
                context['error'] = 'El archivo debe ser .xlsx o .xls'
                return render(request, 'admin/importar_prospectos.html', context)

            wb = openpyxl.load_workbook(archivo, data_only=True)
            ws = wb.active

            creados = 0
            actualizados = 0
            errores = []

            def _normalizar_celda(val):
                if val is None:
                    return ''
                if isinstance(val, (int, float)):
                    return str(int(val)) if isinstance(val, float) and val == int(val) else str(val) if isinstance(val, float) else str(val)
                return str(val).strip()

            def _normalizar_telefono(raw):
                tel = re.sub(r'\D', '', raw)
                if tel.startswith('57') and len(tel) == 12:
                    return tel
                if len(tel) == 10 and tel.startswith('3'):
                    return '57' + tel
                if len(tel) == 7 or len(tel) == 10:
                    return '57' + tel
                return tel

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or all(cell is None or str(cell).strip() == '' for cell in row[:2]):
                    continue

                try:
                    telefono_raw = _normalizar_celda(row[0]) if len(row) > 0 else ''
                    nombre = _normalizar_celda(row[1]) if len(row) > 1 else ''
                    email = _normalizar_celda(row[2]) if len(row) > 2 else ''
                    empresa = _normalizar_celda(row[3]) if len(row) > 3 else ''

                    if not telefono_raw:
                        errores.append(f"Fila {row_idx}: Teléfono vacío")
                        continue

                    telefono = _normalizar_telefono(telefono_raw)
                    if not telefono or len(telefono) < 10:
                        errores.append(f"Fila {row_idx}: Teléfono inválido '{telefono_raw}'")
                        continue

                    from ..models import ProspectoB2B
                    prospecto, created = ProspectoB2B.objects.update_or_create(
                        telefono=telefono,
                        defaults={
                            'nombre_contacto': nombre or '',
                            'email': email or '',
                            'empresa': empresa or '',
                            'origen': 'excel',
                        }
                    )
                    if created:
                        creados += 1
                    else:
                        actualizados += 1

                except Exception as e:
                    errores.append(f"Fila {row_idx}: {str(e)}")

            context.update({
                'exito': True,
                'creados': creados,
                'actualizados': actualizados,
                'total': creados + actualizados,
                'advertencias': errores[:20] if errores else [],
            })
        except Exception as e:
            context['error'] = f"Error procesando archivo: {str(e)}"

    return render(request, 'admin/importar_prospectos.html', context)


# ---------- Vista de importación de estudiantes ----------
@staff_member_required
def importar_estudiantes(request):
    """Compat: redirige al importador LatAm del ModelAdmin (fuente de verdad)."""
    from django.shortcuts import redirect

    return redirect('admin:core_estudiante_importar')



# ---------- Vista de descarga de reportes ----------
@staff_member_required
def descargar_reportes(request):
    """Vista para descargar reportes en Excel filtrando por fechas."""
    context = {}
    
    if request.method == 'POST':
        fecha_inicio = request.POST.get('fecha_inicio')
        fecha_fin = request.POST.get('fecha_fin')
        tipo_reporte = request.POST.get('tipo_reporte', 'todos')  # todos, envios, whatsapp
        
        try:
            # Parsear fechas
            inicio = datetime.strptime(fecha_inicio, '%Y-%m-%d') if fecha_inicio else None
            fin = datetime.strptime(fecha_fin, '%Y-%m-%d') if fecha_fin else None
            
            # Ajustar fin de día
            if fin:
                fin = fin.replace(hour=23, minute=59, second=59)
            
            # Crear workbook
            wb = openpyxl.Workbook()
            wb.remove(wb.active)  # Eliminar hoja por defecto
            
            # Estilos
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            # ========== ENVÍOS ==========
            if tipo_reporte in ['todos', 'envios']:
                ws_envios = wb.create_sheet('Envíos')
                
                # Filtrar por fecha
                queryset = EnvioLog.objects.all()
                if inicio:
                    queryset = queryset.filter(fecha_envio__gte=inicio)
                if fin:
                    queryset = queryset.filter(fecha_envio__lte=fin)
                queryset = queryset.order_by('-fecha_envio')
                
                # Encabezados
                headers = ['ID', 'Estudiante', 'Teléfono', 'Campaña', 'Plantilla', 'Estado', 'Fecha', 'Respuesta API']
                ws_envios.append(headers)
                
                # Aplicar estilos a encabezados
                for cell in ws_envios[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                # Datos
                for log in queryset:
                    fecha_str = log.fecha_envio.strftime('%Y-%m-%d %H:%M:%S') if log.fecha_envio else ''
                    row = [
                        log.id,
                        log.estudiante.nombre,
                        log.estudiante.telefono,
                        log.campana.nombre,
                        log.campana.plantilla.nombre_interno,
                        log.estado,
                        fecha_str,
                        log.respuesta_api or ''
                    ]
                    ws_envios.append(row)
                
                # Ajustar ancho de columnas
                ws_envios.column_dimensions['A'].width = 8
                ws_envios.column_dimensions['B'].width = 20
                ws_envios.column_dimensions['C'].width = 15
                ws_envios.column_dimensions['D'].width = 20
                ws_envios.column_dimensions['E'].width = 20
                ws_envios.column_dimensions['F'].width = 12
                ws_envios.column_dimensions['G'].width = 20
                ws_envios.column_dimensions['H'].width = 30
            
            # ========== WHATSAPP ==========
            if tipo_reporte in ['todos', 'whatsapp']:
                ws_whatsapp = wb.create_sheet('WhatsApp')
                
                # Filtrar por fecha
                queryset = WhatsappLog.objects.all()
                if inicio:
                    queryset = queryset.filter(fecha__gte=inicio)
                if fin:
                    queryset = queryset.filter(fecha__lte=fin)
                queryset = queryset.order_by('-fecha')
                
                # Encabezados
                headers = ['ID', 'Teléfono', 'Tipo', 'Estado', 'Mensaje', 'Fecha', 'ID Mensaje']
                ws_whatsapp.append(headers)
                
                # Aplicar estilos a encabezados
                for cell in ws_whatsapp[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                # Datos
                for log in queryset:
                    fecha_str = log.fecha.strftime('%Y-%m-%d %H:%M:%S') if log.fecha else ''
                    tipo = '📥 Entrante' if log.tipo == 'INCOMING' else '📤 Saliente'
                    row = [
                        log.id,
                        log.telefono,
                        tipo,
                        log.estado,
                        log.mensaje or '',
                        fecha_str,
                        log.mensaje_id or ''
                    ]
                    ws_whatsapp.append(row)
                
                # Ajustar ancho de columnas
                ws_whatsapp.column_dimensions['A'].width = 8
                ws_whatsapp.column_dimensions['B'].width = 15
                ws_whatsapp.column_dimensions['C'].width = 15
                ws_whatsapp.column_dimensions['D'].width = 12
                ws_whatsapp.column_dimensions['E'].width = 50
                ws_whatsapp.column_dimensions['F'].width = 20
                ws_whatsapp.column_dimensions['G'].width = 25
            
            # Generar respuesta
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            fecha_str = datetime.now().strftime('%Y%m%d_%H%M%S')
            response['Content-Disposition'] = f'attachment; filename="Reporte_eki_{fecha_str}.xlsx"'
            wb.save(response)
            return response
        
        except Exception as e:
            context['error'] = f"Error al generar reporte: {str(e)}"
    
    # GET: mostrar formulario
    # Calcular primer día del mes actual y último día
    hoy = datetime.now()
    primer_dia_mes = hoy.replace(day=1)
    if hoy.month == 12:
        ultimo_dia_mes = primer_dia_mes.replace(year=hoy.year + 1, month=1, day=1) - timedelta(days=1)
    else:
        ultimo_dia_mes = primer_dia_mes.replace(month=hoy.month + 1, day=1) - timedelta(days=1)
    
    context['fecha_inicio_default'] = primer_dia_mes.strftime('%Y-%m-%d')
    context['fecha_fin_default'] = ultimo_dia_mes.strftime('%Y-%m-%d')
    
    return render(request, 'admin/descargar_reportes.html', context)


@staff_member_required
def probar_twilio_view(request):
    """Vista para probar integración con Twilio WhatsApp"""
    context = {
        'mensaje': None,
        'error': False,
        'resultado': None
    }
    
    if request.method == 'POST':
        try:
            from twilio.rest import Client
            import os
            
            # Obtener datos del formulario
            tipo_mensaje = request.POST.get('tipo_mensaje')
            usar_template = request.POST.get('usar_template') == 'on'
            telefono = request.POST.get('telefono', '').strip()
            mensaje_texto = request.POST.get('mensaje', '').strip()
            url_imagen = request.POST.get('url_imagen', '').strip()
            
            # Validar credenciales
            account_sid = os.environ.get('TWILIO_ACCOUNT_SID')
            auth_token = os.environ.get('TWILIO_AUTH_TOKEN')
            template_sid = os.environ.get('TWILIO_TEMPLATE_SID')
            
            if not account_sid or not auth_token:
                context['mensaje'] = '<strong>❌ Error:</strong> Las credenciales de Twilio no están configuradas en el archivo .env'
                context['error'] = True
                return render(request, 'admin/probar_twilio.html', context)
            
            # Validar teléfono
            if not telefono:
                context['mensaje'] = '<strong>❌ Error:</strong> Debes proporcionar un número de teléfono'
                context['error'] = True
                return render(request, 'admin/probar_twilio.html', context)
            
            # Asegurar formato whatsapp:
            if not telefono.startswith('+'):
                telefono = f'+{telefono}'
            if not telefono.startswith('whatsapp:'):
                telefono_whatsapp = f'whatsapp:{telefono}'
            else:
                telefono_whatsapp = telefono
            
            # Crear cliente Twilio
            client = Client(account_sid, auth_token)
            
            # Si se usa template aprobado
            if usar_template and template_sid:
                message = client.messages.create(
                    content_sid=template_sid,
                    from_=getattr(settings, 'TWILIO_PHONE_NUMBER', 'whatsapp:+573202948806'),
                    to=telefono_whatsapp
                )
            else:
                # Preparar parámetros del mensaje libre
                params = {
                    "to": telefono_whatsapp,
                    "from_": getattr(settings, 'TWILIO_PHONE_NUMBER', 'whatsapp:+573202948806'),
                    "body": mensaje_texto
                }
                
                # Si es mensaje con imagen/video, generar URL firmada y agregar media_url
                if tipo_mensaje == 'imagen' and url_imagen:
                    # Revisar si la URL es de S3 y necesita firma
                    from core.utils import generar_url_firmada_s3_v4
                    import re
                    s3_pattern = r'https://([\w\-]+)\.s3[\w\-\.]*\.amazonaws\.com/(.+)'
                    match = re.match(s3_pattern, url_imagen)
                    if match:
                        bucket_name = match.group(1)
                        object_name = match.group(2)
                        url_firmada = generar_url_firmada_s3_v4(bucket_name, object_name)
                        params["media_url"] = [url_firmada]
                    else:
                        params["media_url"] = [url_imagen]
                
                # Enviar mensaje
                message = client.messages.create(**params)
            
            # Crear resultado formateado
            resultado_texto = f"""
✅ MENSAJE ENVIADO EXITOSAMENTE

📝 SID: {message.sid}
📊 Estado: {message.status}
📅 Fecha: {message.date_created}
📱 Destino: {telefono}
"""
            
            if usar_template and template_sid:
                resultado_texto += f"📋 Template SID: {template_sid}\n"
            else:
                resultado_texto += f"💬 Mensaje: {mensaje_texto[:100]}{'...' if len(mensaje_texto) > 100 else ''}\n"
                if tipo_mensaje == 'imagen' and url_imagen:
                    resultado_texto += f"🖼️  Imagen: {url_imagen}\n"
            
            context['mensaje'] = f'<strong>✅ ¡Éxito!</strong> El mensaje fue enviado correctamente. SID: {message.sid}'
            context['error'] = False
            context['resultado'] = resultado_texto
            
            # Guardar log
            WhatsappLog.objects.create(
                telefono=telefono.replace('whatsapp:', '').replace('+', ''),
                mensaje=mensaje_texto,
                mensaje_id=message.sid,
                estado='SENT'
            )
            
        except Exception as e:
            context['mensaje'] = f'<strong>❌ Error al enviar:</strong> {str(e)}'
            context['error'] = True
            context['resultado'] = f"ERROR:\n{str(e)}"
    
    return render(request, 'admin/probar_twilio.html', context)


@staff_member_required
def calendario_campanas_view(request):
    """Vista de calendario de campañas programadas"""
    from django.utils import timezone
    
    ahora = timezone.now()
    
    # Campañas pendientes (programadas pero no ejecutadas)
    campanas_pendientes = Campana.objects.filter(
        fecha_programada__isnull=False,
        ejecutada=False
    ).order_by('fecha_programada')
    
    # Campañas ejecutadas que tenían programación
    campanas_ejecutadas = Campana.objects.filter(
        fecha_programada__isnull=False,
        ejecutada=True
    ).order_by('-fecha_programada')[:10]
    
    context = {
        'campanas_pendientes': campanas_pendientes,
        'campanas_ejecutadas': campanas_ejecutadas,
    }
    
    return render(request, 'admin/calendario_campanas.html', context)


@staff_member_required
def conversaciones_view(request):
    """Vista de conversaciones estilo WhatsApp Web (pantalla completa)."""
    from core.conversaciones_service import construir_contexto_inbox

    cliente_filtro_raw = (request.GET.get("cliente") or "").strip()
    cliente_filtro_id = int(cliente_filtro_raw) if cliente_filtro_raw.isdigit() else None
    estudiante_raw = (request.GET.get("estudiante") or "").strip()
    estudiante_id = int(estudiante_raw) if estudiante_raw.isdigit() else None
    page_raw = (request.GET.get("page") or "1").strip()
    page = int(page_raw) if page_raw.isdigit() else 1

    context = construir_contexto_inbox(
        cliente_filtro_id=cliente_filtro_id,
        estudiante_id=estudiante_id,
        telefono=(request.GET.get("telefono") or "").strip() or None,
        busqueda=(request.GET.get("q") or "").strip() or None,
        page=page,
    )
    context.update({
        "inbox_base_url": "/admin/conversaciones/",
        "inbox_volver_url": "/admin/",
        "inbox_volver_label": "Panel admin",
        "inbox_modo": "admin",
    })
    if estudiante_id:
        from django.urls import reverse

        context["inbox_volver_url"] = reverse(
            "admin:core_estudiante_change",
            args=[estudiante_id],
        )
        context["inbox_volver_label"] = "Ficha estudiante"
    elif cliente_filtro_id:
        from django.urls import reverse

        context["inbox_volver_url"] = reverse(
            "admin:core_cliente_change",
            args=[cliente_filtro_id],
        )
        context["inbox_volver_label"] = "Ficha cliente"
    return render(request, "admin/conversaciones.html", context)


@staff_member_required
def chat_prueba_view(request):
    """Vista para probar la IA sin necesidad de WhatsApp/ngrok"""
    return render(request, 'admin/chat_prueba.html')


@staff_member_required
def chat_prueba_api(request):
    """API para el chat de prueba"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            mensaje = data.get('mensaje', '')
            telefono = data.get('telefono', 'test_chat')
            
            print(f"🔵 Chat de prueba - Mensaje: {mensaje}")
            
            # Guardar mensaje entrante
            WhatsappLog.objects.create(
                telefono=telefono,
                mensaje=mensaje,
                mensaje_id=f"test_{timezone.now().timestamp()}",
                tipo='INCOMING'
            )
            
            # Obtener respuesta de la IA
            try:
                from ..ai_assistant import responder_con_ia
                respuesta = responder_con_ia(mensaje, telefono)
                print(f"✅ IA respondió: {respuesta}")
            except Exception as e:
                print(f"❌ Error en IA: {e}")
                # Fallback
                from ..intent_detector import detect_intent
                from ..response_templates import get_response_for_intent
                intent = detect_intent(mensaje)
                respuesta = get_response_for_intent(intent, 'Usuario')
            
            # Guardar respuesta
            WhatsappLog.objects.create(
                telefono=telefono,
                mensaje=respuesta,
                mensaje_id=f"test_response_{timezone.now().timestamp()}",
                tipo='SENT'
            )
            
            return JsonResponse({
                'success': True,
                'respuesta': respuesta
            })
            
        except Exception as e:
            print(f"❌ Error en chat de prueba: {e}")
            return JsonResponse({
                'success': False,
                'error': str(e)
            }, status=500)
    
    return JsonResponse({'error': 'Método no permitido'}, status=405)


@staff_member_required
def test_email_gmail_view(request):
    """Vista para probar la configuración de Gmail"""
    from ..email_test import test_gmail_connection, format_email_status_html
    
    context = {
        'title': 'Probar Conexión Gmail',
        'status_html': format_email_status_html(),
        'resultado': None
    }
    
    if request.method == 'POST':
        success, message = test_gmail_connection()
        context['resultado'] = {
            'success': success,
            'message': message
        }
    
    return render(request, 'admin/test_email.html', context)


# ========================================
# VISTAS PARA GENERACIÓN DE CURSOS CON IA
# ========================================

def _contexto_crear_curso_ia(request):
    from ..models import Cliente
    from ..utils_ia import MODELO_IA_DEFAULT, modelos_ia_habilitados

    historial = request.session.get('chat_curso_ia') or []
    modelos = modelos_ia_habilitados()
    return {
        'clientes': Cliente.objects.filter(activo=True).order_by('nombre'),
        'modelos_ia': modelos,
        'modelos_ia_disponibles': bool(modelos),
        'historial_chat': historial,
        'modelo_actual': request.session.get('modelo_usado', MODELO_IA_DEFAULT),
    }


def _guardar_sesion_curso_ia(request, estructura, cliente_id, fuente_nombre, modelo_ia, texto_fuente=''):
    request.session['estructura_curso'] = estructura
    request.session['cliente_id'] = str(cliente_id)
    request.session['archivo_nombre'] = fuente_nombre or 'prompt-ia'
    request.session['modelo_usado'] = modelo_ia
    request.session['texto_fuente_curso'] = (texto_fuente or '')[:12000]
    request.session.modified = True


@staff_member_required
def subir_documento_curso(request):
    """Paso 1: chat + prompt largo o archivo → estructura JSON del curso."""
    from django.shortcuts import redirect
    from ..models import Cliente
    from ..utils_ia import (
        extraer_texto_documento,
        generar_estructura_curso_con_ia,
        validar_estructura_curso,
    )

    context = _contexto_crear_curso_ia(request)

    if request.method == 'POST':
        try:
            accion = request.POST.get('accion', 'generar')
            cliente_id = request.POST.get('cliente_id')
            modelo_ia = request.POST.get('modelo_ia', 'gpt-4o-mini')
            prompt_usuario = (request.POST.get('prompt') or '').strip()
            archivo = request.FILES.get('documento')

            if accion == 'limpiar_chat':
                request.session.pop('chat_curso_ia', None)
                return redirect('subir_documento_curso')

            if not cliente_id:
                context['error'] = 'Selecciona una organización'
                return render(request, 'admin/subir_documento_curso.html', context)

            try:
                cliente = Cliente.objects.get(id=cliente_id)
            except Cliente.DoesNotExist:
                context['error'] = 'Organización no encontrada'
                return render(request, 'admin/subir_documento_curso.html', context)

            texto = ''
            fuente = 'prompt-ia'
            if archivo:
                nombre = archivo.name.lower()
                if not (nombre.endswith('.pdf') or nombre.endswith('.docx') or nombre.endswith('.txt')):
                    context['error'] = 'Solo PDF, Word (.docx) o TXT'
                    return render(request, 'admin/subir_documento_curso.html', context)
                texto = extraer_texto_documento(archivo)
                fuente = archivo.name
            elif prompt_usuario:
                texto = prompt_usuario
                fuente = 'prompt-chat'
            else:
                context['error'] = 'Escribe un prompt o sube un documento'
                return render(request, 'admin/subir_documento_curso.html', context)

            if len(texto) < 200:
                context['error'] = 'El contenido es muy corto (mínimo 200 caracteres)'
                return render(request, 'admin/subir_documento_curso.html', context)

            from ..utils_ia import validar_modelo_ia_disponible
            try:
                validar_modelo_ia_disponible(modelo_ia)
            except ValueError as e:
                context['error'] = str(e)
                return render(request, 'admin/subir_documento_curso.html', context)

            historial = list(request.session.get('chat_curso_ia') or [])
            historial.append({'rol': 'user', 'texto': prompt_usuario or f'[Archivo: {fuente}]'})
            request.session['chat_curso_ia'] = historial[-20:]

            use_async = not getattr(settings, 'CELERY_TASK_ALWAYS_EAGER', False)
            if use_async:
                import uuid
                from django.core.cache import cache
                from core.tasks import generar_curso_ia_async, _curso_ia_cache_key

                job_id = str(uuid.uuid4())
                cache.set(_curso_ia_cache_key(job_id), {'status': 'pending'}, 3600)
                request.session['curso_ia_job_id'] = job_id
                request.session['curso_ia_pending'] = {
                    'cliente_id': str(cliente_id),
                    'fuente': fuente,
                    'modelo_ia': modelo_ia,
                    'texto': texto[:12000],
                    'prompt_usuario': prompt_usuario,
                }
                request.session.modified = True
                try:
                    generar_curso_ia_async.delay(job_id, texto, modelo_ia)
                    return redirect('generando_curso_ia')
                except Exception as celery_err:
                    logger.warning('Celery no disponible para curso IA, modo sync: %s', celery_err)
                    request.session.pop('curso_ia_job_id', None)
                    request.session.pop('curso_ia_pending', None)

            estructura = generar_estructura_curso_con_ia(texto, modelo=modelo_ia)
            es_valida, errores = validar_estructura_curso(estructura)
            if not es_valida:
                context['error'] = f'Estructura inválida: {", ".join(errores)}'
                return render(request, 'admin/subir_documento_curso.html', context)

            historial.append({
                'rol': 'assistant',
                'texto': f'Generé «{estructura.get("titulo", "Curso")}» con {len(estructura.get("modulos", []))} módulos.',
            })
            request.session['chat_curso_ia'] = historial[-20:]
            _guardar_sesion_curso_ia(request, estructura, cliente_id, fuente, modelo_ia, texto)
            return redirect('vista_previa_curso_ia')

        except ValueError as e:
            context['error'] = str(e)
        except Exception as e:
            context['error'] = f'Error: {e}'
            logger.error(f'Error en subir_documento_curso: {e}', exc_info=True)

    return render(request, 'admin/subir_documento_curso.html', context)


@staff_member_required
def generando_curso_ia(request):
    """Pantalla de espera mientras Celery genera la estructura (evita 504)."""
    job_id = request.session.get('curso_ia_job_id')
    if not job_id:
        return redirect('subir_documento_curso')
    return render(request, 'admin/generando_curso_ia.html', {'job_id': job_id})


@staff_member_required
def api_estado_curso_ia(request):
    """Polling JSON del job de generación IA."""
    from django.core.cache import cache
    from django.http import JsonResponse
    from core.tasks import _curso_ia_cache_key

    job_id = request.GET.get('job_id') or request.session.get('curso_ia_job_id')
    if not job_id:
        return JsonResponse({'status': 'missing'}, status=404)
    data = cache.get(_curso_ia_cache_key(job_id)) or {'status': 'pending'}
    if data.get('status') == 'ok':
        pending = request.session.get('curso_ia_pending') or {}
        estructura = data.get('estructura')
        if estructura and pending.get('cliente_id'):
            historial = list(request.session.get('chat_curso_ia') or [])
            historial.append({
                'rol': 'assistant',
                'texto': f'Generé «{estructura.get("titulo", "Curso")}» con {len(estructura.get("modulos", []))} módulos.',
            })
            request.session['chat_curso_ia'] = historial[-20:]
            _guardar_sesion_curso_ia(
                request,
                estructura,
                pending['cliente_id'],
                pending.get('fuente', 'prompt-ia'),
                pending.get('modelo_ia', 'gpt-4o-mini'),
                pending.get('texto', ''),
            )
            request.session.pop('curso_ia_job_id', None)
            request.session.pop('curso_ia_pending', None)
            return JsonResponse({
                'status': 'ok',
                'redirect': '/admin/vista-previa-curso-ia/',
                'titulo': estructura.get('titulo'),
                'modulos': len(estructura.get('modulos', [])),
            })
    if data.get('status') == 'error':
        request.session.pop('curso_ia_job_id', None)
        request.session.pop('curso_ia_pending', None)
    return JsonResponse({
        'status': data.get('status', 'pending'),
        'error': data.get('error'),
    })


@staff_member_required
def vista_previa_curso_ia(request):
    """Paso 2: revisión humana, edición y regeneración por módulo."""
    from django.contrib import messages
    from django.shortcuts import redirect
    from ..models import Cliente
    from ..utils_ia import guardar_curso_desde_estructura, regenerar_modulo_en_estructura

    estructura = request.session.get('estructura_curso')
    cliente_id = request.session.get('cliente_id')
    archivo_nombre = request.session.get('archivo_nombre', 'prompt-ia')
    modelo_usado = request.session.get('modelo_usado', 'gpt-4o-mini')
    texto_fuente = request.session.get('texto_fuente_curso', '')

    if not estructura or not cliente_id:
        messages.error(request, 'No hay borrador de curso. Genera uno primero.')
        return redirect('subir_documento_curso')

    try:
        cliente = Cliente.objects.get(id=cliente_id)
    except Cliente.DoesNotExist:
        messages.error(request, 'Organización no encontrada')
        return redirect('subir_documento_curso')

    context = {
        'estructura': estructura,
        'estructura_json': json.dumps(estructura, ensure_ascii=False, indent=2),
        'cliente': cliente,
        'archivo_nombre': archivo_nombre,
        'modelo_usado': modelo_usado,
        'total_modulos': len(estructura.get('modulos', [])),
        'total_lecciones': sum(len(m.get('lecciones', [])) for m in estructura.get('modulos', [])),
        'historial_chat': request.session.get('chat_curso_ia') or [],
    }

    if request.method == 'POST':
        accion = request.POST.get('accion')

        if accion == 'guardar':
            try:
                estructura['titulo'] = request.POST.get('titulo', estructura['titulo'])
                estructura['descripcion'] = request.POST.get('descripcion', estructura['descripcion'])
                estructura['duracion_estimada'] = request.POST.get(
                    'duracion_estimada', estructura.get('duracion_estimada', '4 semanas')
                )
                estructura['nivel'] = request.POST.get('nivel', estructura.get('nivel', 'Intermedio'))
                estructura['puntos_por_leccion'] = int(
                    request.POST.get('puntos_por_leccion', estructura.get('puntos_por_leccion', 50))
                )
                curso = guardar_curso_desde_estructura(estructura, cliente, archivo_nombre)
                for key in (
                    'estructura_curso', 'cliente_id', 'archivo_nombre',
                    'modelo_usado', 'texto_fuente_curso', 'chat_curso_ia',
                ):
                    request.session.pop(key, None)
                messages.success(
                    request,
                    f'Curso «{curso.nombre}» creado (inactivo). Revísalo en el admin antes de activar.',
                )
                return redirect(f'/admin/core/curso/{curso.id}/change/')
            except Exception as e:
                context['error'] = f'Error al guardar: {e}'
                logger.error(f'Error guardando curso IA: {e}', exc_info=True)

        elif accion == 'regenerar_modulo':
            try:
                idx = int(request.POST.get('modulo_indice', 0))
                instrucciones = (request.POST.get('instrucciones_regenerar') or '').strip()
                estructura = regenerar_modulo_en_estructura(
                    estructura,
                    idx,
                    texto_fuente=texto_fuente,
                    instrucciones=instrucciones,
                    modelo=modelo_usado,
                )
                request.session['estructura_curso'] = estructura
                historial = list(request.session.get('chat_curso_ia') or [])
                historial.append({
                    'rol': 'assistant',
                    'texto': f'Regeneré el módulo {idx + 1}: {estructura["modulos"][idx].get("nombre", "")}',
                })
                request.session['chat_curso_ia'] = historial[-20:]
                messages.success(request, f'Módulo {idx + 1} regenerado.')
                return redirect('vista_previa_curso_ia')
            except Exception as e:
                context['error'] = f'No se pudo regenerar: {e}'

        elif accion == 'cancelar':
            for key in (
                'estructura_curso', 'cliente_id', 'archivo_nombre',
                'modelo_usado', 'texto_fuente_curso', 'chat_curso_ia',
            ):
                request.session.pop(key, None)
            messages.info(request, 'Creación cancelada')
            return redirect('subir_documento_curso')

        context['estructura'] = estructura
        context['estructura_json'] = json.dumps(estructura, ensure_ascii=False, indent=2)
        context['total_modulos'] = len(estructura.get('modulos', []))

    return render(request, 'admin/vista_previa_curso_ia.html', context)
